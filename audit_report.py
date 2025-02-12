import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from datetime import datetime


def write_header(ws, vets):
    """
    Writes the header row and transcript row to the Excel sheet.
    """
    # Title row
    title_fill = PatternFill(start_color="0000FF", end_color="0000FF", fill_type="solid")
    title_font = Font(size=14, bold=True, color="FFFFFF")

    ws.merge_cells('A1:D1')
    ws['A1'] = "Audit Scoreboard"
    ws['A1'].fill = title_fill
    ws['A1'].font = title_font

    # Header row
    header_row = ["Parameters", "Weights", "Subparameters", "Subweights"] + [vet['name'] for vet in vets]
    ws.append(header_row)

    # Bold header row
    for col_num, header in enumerate(header_row, start=1):
        ws.cell(row=2, column=col_num).font = Font(size=12, bold=True)

    # Transcript row
    transcript_row = ["", "", "", ""] + [vet.get('transcript_file', "N/A") for vet in vets]
    ws.append(transcript_row)


def write_total_row(ws, current_row, vets, subparam_list, start_row):
    """
    Writes the total row for a parameter.
    """
    total_row = ["", "", "Total", ""]
    bold_font = Font(bold=True)  # Define a bold font style

    for vet_index in range(len(vets)):
        if subparam_list:  # If subparameters exist
            start_cell = ws.cell(row=start_row, column=5 + vet_index).coordinate
            end_cell = ws.cell(row=current_row - 1, column=5 + vet_index).coordinate
            total_row.append(f"=SUM({start_cell}:{end_cell})")
        else:
            total_row.append("")

    ws.append(total_row)

    # Apply bold font to the 'Total' row
    for col_num in range(1, len(total_row) + 1):  # Iterate through all columns in the total row
        ws.cell(row=ws.max_row, column=col_num).font = bold_font


def write_parameter_data(ws, parameters, weights, subparameters, subweights, scores, vets, start_row):
    """
    Writes the parameters, subparameters, and scores to the Excel sheet.
    """
    current_row = start_row
    italic_font = Font(italic=True)  # Define an italic font style

    for param_index, (param, weight, subparam_list, subweight_list) in enumerate(zip(parameters, weights, subparameters, subweights)):
        # Write parameter name and weight
        ws.cell(row=current_row, column=1, value=param)
        ws.cell(row=current_row, column=2, value=weight)

        if subparam_list:  # If there are subparameters
            for subparam_index, (subparam, subweight) in enumerate(zip(subparam_list, subweight_list)):

                # Write the subparameter name and apply italic style
                subparam_cell = ws.cell(row=current_row, column=3, value=subparam)
                subparam_cell.font = italic_font

                # Write subweight
                ws.cell(row=current_row, column=4, value=subweight)

                for vet_index, vet in enumerate(vets):
                    # Access subparameter-specific score
                    score = scores[param_index][vet_index]["subscores"][subparam_index]
                    ws.cell(row=current_row, column=5 + vet_index, value=score)

                current_row += 1
        else:  # If there are no subparameters
            for vet_index, vet in enumerate(vets):
                score = scores[param_index][vet_index]["score"]
                ws.cell(row=current_row, column=5 + vet_index, value=score)

            current_row += 1

        # Write total row for the parameter
        write_total_row(ws, current_row, vets, subparam_list, current_row - len(subparam_list) if subparam_list else current_row - 1)
        current_row += 1

    return current_row


def generate_audit_excel_report(vets, parameters, weights, subparameters, subweights, scores):
    """
    Generates an audit Excel report and saves it to the output directory.
    """
    # Create a new workbook and select the active worksheet
    wb = Workbook()
    ws = wb.active

    # Write the header rows
    write_header(ws, vets)

    # Write the parameter data
    start_row = 4
    write_parameter_data(ws, parameters, weights, subparameters, subweights, scores, vets, start_row)

    # Ensure the output directory exists
    output_dir = "./output"
    os.makedirs(output_dir, exist_ok=True)

    # Save the workbook with a timestamped filename
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_filename = f"{output_dir}/audit_scoreboard_{timestamp}.xlsx"
    wb.save(output_filename)
    print(f"Audit report saved as: {output_filename}")